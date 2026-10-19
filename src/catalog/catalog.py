"""Import and validate the authoritative ITEMCODE_SKU workbook.

``item_code`` is always the identity.  Workbook row numbers and generated
folder labels are traceability/display fields only.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable


REQUIRED_COLUMNS = ("STT", "CATEGORY", "BRAND", "ITEM_CODE", "PRODUCT_NAME_VN")
OPTIONAL_LEGACY_COLUMN = "CLASS LABEL"
WINDOWS_RESERVED = {
    "CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}
PACKAGING_RE = re.compile(
    r"(?i)(\d+(?:[.,]\d+)?\s*(?:ml|kg|g|l)"
    r"(?:\s*(?:[/x×-]\s*)?\d+\s*(?:chai|gói|goi|hộp|hop|tuýp|tuyp|túi|tui|thùng|thung))?)"
)


@dataclass(frozen=True)
class CatalogRecord:
    item_code: str
    item_name: str
    brand: str
    category: str
    class_name: str
    filesystem_name: str
    workbook_reference: dict[str, Any]
    packaging: str | None = None
    legacy_class_name: str | None = None
    duplicate_workbook_references: tuple[dict[str, Any], ...] = field(default_factory=tuple)

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["duplicate_workbook_references"] = list(self.duplicate_workbook_references)
        return value


@dataclass(frozen=True)
class CatalogIssue:
    severity: str
    code: str
    message: str
    workbook_reference: str | None = None
    item_code: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _ascii_label(value: str) -> str:
    value = value.replace("Đ", "D").replace("đ", "d")
    value = "".join(
        char for char in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(char)
    )
    pieces = re.findall(r"[A-Za-z0-9]+", value)
    return "_".join(pieces) or "unnamed_product"


def filesystem_safe_name(value: str, max_length: int = 180) -> str:
    """Return a readable deterministic Windows-safe folder name."""

    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value).strip(" ._")
    safe = re.sub(r"_+", "_", safe) or "unnamed_product"
    if safe.upper() in WINDOWS_RESERVED:
        safe = f"_{safe}"
    if len(safe) > max_length:
        suffix = hashlib.sha256(safe.encode("utf-8")).hexdigest()[:10]
        safe = f"{safe[: max_length - 11].rstrip(' ._')}_{suffix}"
    return safe


def extract_packaging(item_name: str) -> str | None:
    """Extract packaging text verbatim; the original name remains authoritative."""

    matches = list(PACKAGING_RE.finditer(item_name))
    return matches[-1].group(1).strip() if matches else None


def _numeric_item_code(cell: Any, value: int | float) -> tuple[str, CatalogIssue | None]:
    if isinstance(value, float) and not value.is_integer():
        return "", CatalogIssue(
            "error", "non_integral_numeric_item_code",
            f"ITEM_CODE numeric value is not an integer: {value!r}", cell.coordinate,
        )
    digits = str(int(value))
    number_format = str(cell.number_format or "General")
    # A zero-only Excel format explicitly carries a fixed width and can safely
    # restore leading zeroes.  General/0 cannot prove that none were lost.
    if re.fullmatch(r"0{2,}", number_format):
        return digits.zfill(len(number_format)), None
    return digits, CatalogIssue(
        "warning", "ambiguous_numeric_item_code_format",
        (
            f"ITEM_CODE is stored as numeric ({digits}) with format {number_format!r}; "
            "leading-zero intent cannot be proven. The numeric value is reported explicitly."
        ),
        cell.coordinate, digits,
    )


def item_code_from_cell(cell: Any) -> tuple[str, CatalogIssue | None]:
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        return "", CatalogIssue("error", "missing_item_code", "ITEM_CODE is empty", cell.coordinate)
    if isinstance(value, bool):
        return "", CatalogIssue("error", "invalid_item_code", "Boolean ITEM_CODE is invalid", cell.coordinate)
    if isinstance(value, (int, float)):
        return _numeric_item_code(cell, value)
    text = str(value).strip()
    if text.startswith("="):
        return "", CatalogIssue(
            "error", "formula_item_code",
            "ITEM_CODE is a formula; replace it with an explicit text identifier", cell.coordinate,
        )
    return text, None


def validate_records(records: Iterable[CatalogRecord]) -> list[CatalogIssue]:
    """Validate cross-record conflicts, folder collisions, and legacy ambiguity."""

    records = list(records)
    issues: list[CatalogIssue] = []
    by_code: dict[str, list[CatalogRecord]] = defaultdict(list)
    by_folder: dict[str, list[CatalogRecord]] = defaultdict(list)
    by_legacy: dict[str, list[CatalogRecord]] = defaultdict(list)
    for record in records:
        by_code[record.item_code].append(record)
        by_folder[record.filesystem_name.casefold()].append(record)
        if record.legacy_class_name:
            by_legacy[record.legacy_class_name.casefold()].append(record)
    for item_code, group in by_code.items():
        identities = {(x.item_name, x.brand, x.category, x.packaging) for x in group}
        if len(identities) > 1:
            refs = ", ".join(x.workbook_reference["cell"] for x in group)
            issues.append(CatalogIssue(
                "error", "conflicting_item_code",
                f"Item code {item_code} has conflicting records at {refs}", refs, item_code,
            ))
    for group in by_folder.values():
        codes = {x.item_code for x in group}
        if len(codes) > 1:
            issues.append(CatalogIssue(
                "error", "filesystem_name_collision",
                f"Filesystem name {group[0].filesystem_name!r} maps to item codes {sorted(codes)}",
            ))
    for group in by_legacy.values():
        codes = {x.item_code for x in group}
        if len(codes) > 1:
            issues.append(CatalogIssue(
                "error", "ambiguous_legacy_class",
                f"Legacy class {group[0].legacy_class_name!r} maps to item codes {sorted(codes)}",
            ))
    return issues


def import_workbook(workbook_path: str | Path, sheet_name: str | None = None) -> dict[str, Any]:
    """Read workbook values without guessing identities from names or colors."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised by CLI environment
        raise RuntimeError("Catalog import requires openpyxl; install requirements.txt") from exc

    path = Path(workbook_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Authoritative workbook not found: {path}")
    workbook = load_workbook(path, read_only=False, data_only=False)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1))
    headers = {str(cell.value).strip(): index for index, cell in enumerate(header_cells) if cell.value is not None}
    missing_columns = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing_columns:
        raise ValueError(f"Workbook is missing required columns: {missing_columns}")

    raw_records: list[CatalogRecord] = []
    invalid_rows: list[dict[str, Any]] = []
    issues: list[CatalogIssue] = []
    numeric_item_code_rows = 0
    for row in worksheet.iter_rows(min_row=2):
        relevant = [row[headers[name]].value for name in REQUIRED_COLUMNS]
        if all(value is None or str(value).strip() == "" for value in relevant):
            continue
        code_cell = row[headers["ITEM_CODE"]]
        item_code, code_issue = item_code_from_cell(code_cell)
        if isinstance(code_cell.value, (int, float)) and not isinstance(code_cell.value, bool):
            numeric_item_code_rows += 1
        if code_issue:
            issues.append(code_issue)
        name_cell = row[headers["PRODUCT_NAME_VN"]]
        item_name = "" if name_cell.value is None else str(name_cell.value).strip()
        if not item_name:
            issues.append(CatalogIssue(
                "error", "missing_item_name", "PRODUCT_NAME_VN is empty", name_cell.coordinate, item_code or None,
            ))
        if not item_code or not item_name:
            invalid_rows.append({
                "workbook_reference": {
                    "workbook": path.name,
                    "sheet": worksheet.title,
                    "row": int(code_cell.row),
                    "item_code_cell": code_cell.coordinate,
                    "item_name_cell": name_cell.coordinate,
                },
                "row_values": {
                    "STT": row[headers["STT"]].value,
                    "CATEGORY": row[headers["CATEGORY"]].value,
                    "BRAND": row[headers["BRAND"]].value,
                    "ITEM_CODE": code_cell.value,
                    "PRODUCT_NAME_VN": name_cell.value,
                    "CLASS LABEL": (
                        row[headers[OPTIONAL_LEGACY_COLUMN]].value
                        if OPTIONAL_LEGACY_COLUMN in headers else None
                    ),
                },
                "reason": "missing_item_code" if not item_code else "missing_item_name",
            })
            continue
        brand = "" if row[headers["BRAND"]].value is None else str(row[headers["BRAND"]].value).strip()
        category = "" if row[headers["CATEGORY"]].value is None else str(row[headers["CATEGORY"]].value).strip()
        legacy = None
        if OPTIONAL_LEGACY_COLUMN in headers:
            legacy_value = row[headers[OPTIONAL_LEGACY_COLUMN]].value
            if legacy_value is not None and str(legacy_value).strip() not in {"", "0"}:
                legacy = str(legacy_value).strip()
        label = _ascii_label(legacy or item_name)
        class_name = f"{item_code}_{label}"
        reference = {
            "workbook": path.name,
            "sheet": worksheet.title,
            "row": int(code_cell.row),
            "cell": code_cell.coordinate,
            "stt": row[headers["STT"]].value,
            "item_code_cell_type": code_cell.data_type,
            "item_code_number_format": code_cell.number_format,
        }
        raw_records.append(CatalogRecord(
            item_code=item_code,
            item_name=item_name,
            brand=brand,
            category=category,
            class_name=class_name,
            filesystem_name=filesystem_safe_name(class_name),
            workbook_reference=reference,
            packaging=extract_packaging(item_name),
            legacy_class_name=legacy,
        ))

    # Preserve one record for byte-identical duplicates and make all source rows traceable.
    deduplicated: list[CatalogRecord] = []
    exact_groups: dict[tuple[str, str, str, str, str | None], list[CatalogRecord]] = defaultdict(list)
    for record in raw_records:
        exact_groups[(record.item_code, record.item_name, record.brand, record.category, record.legacy_class_name)].append(record)
    for group in exact_groups.values():
        first = group[0]
        if len(group) > 1:
            refs = tuple(item.workbook_reference for item in group[1:])
            issues.append(CatalogIssue(
                "warning", "duplicate_identical_record",
                f"Identical item record appears {len(group)} times; retained once with all references",
                ", ".join(x["cell"] for x in (first.workbook_reference, *refs)), first.item_code,
            ))
            first = CatalogRecord(**{
                **first.to_dict(), "duplicate_workbook_references": refs,
            })
        deduplicated.append(first)

    issues.extend(validate_records(deduplicated))
    serialized = [record.to_dict() for record in deduplicated]
    catalog_version = "sha256:" + hashlib.sha256(
        json.dumps(serialized, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    errors = sum(issue.severity == "error" for issue in issues)
    warnings = sum(issue.severity == "warning" for issue in issues)
    return {
        "schema_version": 1,
        "catalog_version": catalog_version,
        "source_workbook": str(path),
        "source_workbook_fingerprint": "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_sheet": worksheet.title,
        "status": "invalid" if errors else "valid_with_warnings" if warnings else "valid",
        "record_count": len(deduplicated),
        "raw_record_count": len(raw_records),
        "numeric_item_code_rows": numeric_item_code_rows,
        "records": serialized,
        "invalid_rows": invalid_rows,
        "issues": [issue.to_dict() for issue in issues],
        "summary": {"errors": errors, "warnings": warnings},
    }


def create_scoped_catalog(
    full_bundle: dict[str, Any],
    item_codes: Iterable[str],
    scope_name: str,
) -> dict[str, Any]:
    """Create an explicit allowlist catalog without changing full-catalog validity.

    Every requested item code must resolve to exactly one identity. Byte-identical
    duplicate workbook occurrences are reported; conflicting identities block the
    scoped catalog instead of selecting a row arbitrarily.
    """

    requested = [str(code).strip() for code in item_codes]
    if not requested or any(not code for code in requested):
        raise ValueError("A scoped catalog needs at least one non-empty item code")
    if len(requested) != len(set(requested)):
        raise ValueError("Scoped item codes must be unique and ordered explicitly")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", scope_name):
        raise ValueError("scope_name must use 1-64 letters, digits, '_' or '-'")

    by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in full_bundle["records"]:
        by_code[str(record["item_code"])].append(record)

    selected: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    occurrence_validation: list[dict[str, Any]] = []
    full_issues = list(full_bundle.get("issues", []))
    for code in requested:
        candidates = by_code.get(code, [])
        if not candidates:
            issues.append(CatalogIssue(
                "error", "scoped_item_code_missing",
                f"Requested scoped item code {code} is absent from the workbook", item_code=code,
            ).to_dict())
            occurrence_validation.append({
                "item_code": code, "identity_count": 0, "workbook_occurrence_count": 0,
                "status": "missing",
            })
            continue
        if len(candidates) != 1:
            references = [row["workbook_reference"]["cell"] for row in candidates]
            issues.append(CatalogIssue(
                "error", "scoped_item_code_conflict",
                f"Requested item code {code} has {len(candidates)} conflicting identities at {references}",
                ", ".join(references), code,
            ).to_dict())
            occurrence_validation.append({
                "item_code": code, "identity_count": len(candidates),
                "workbook_occurrence_count": len(candidates), "status": "conflicting",
            })
            continue
        record = candidates[0]
        duplicate_refs = list(record.get("duplicate_workbook_references", []))
        occurrence_count = 1 + len(duplicate_refs)
        occurrence_validation.append({
            "item_code": code,
            "identity_count": 1,
            "workbook_occurrence_count": occurrence_count,
            "workbook_references": [record["workbook_reference"], *duplicate_refs],
            "status": "unique_identity" if occurrence_count == 1 else "identical_duplicates_reported",
        })
        selected.append(record)

    # Carry every issue tied to a selected code, including numeric-format and
    # identical-duplicate warnings. Any scoped error remains blocking.
    issues.extend(issue for issue in full_issues if issue.get("item_code") in set(requested))
    selected_records = [CatalogRecord(**row) for row in selected]
    issues.extend(issue.to_dict() for issue in validate_records(selected_records))
    errors = sum(issue.get("severity") == "error" for issue in issues)
    warnings = sum(issue.get("severity") == "warning" for issue in issues)
    invalid_rows_by_cell = {
        row["workbook_reference"].get("item_code_cell"): row
        for row in full_bundle.get("invalid_rows", [])
    }
    known_out_of_scope = []
    for issue in full_issues:
        if issue.get("severity") != "error" or issue.get("item_code") in set(requested):
            continue
        enriched = dict(issue)
        if issue.get("workbook_reference") in invalid_rows_by_cell:
            enriched["source_row"] = invalid_rows_by_cell[issue["workbook_reference"]]
        known_out_of_scope.append(enriched)
    source_fingerprint = full_bundle.get("source_workbook_fingerprint")
    if not source_fingerprint:
        source_fingerprint = "sha256:" + hashlib.sha256(
            Path(full_bundle["source_workbook"]).read_bytes()
        ).hexdigest()
    version_payload = {
        "schema_version": 1,
        "scope_name": scope_name,
        "selected_item_codes": requested,
        "source_workbook_fingerprint": source_fingerprint,
        "records": selected,
    }
    catalog_version = "sha256:" + hashlib.sha256(
        json.dumps(version_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {
        "schema_version": 1,
        "catalog_version": catalog_version,
        "source_workbook": full_bundle["source_workbook"],
        "source_workbook_fingerprint": source_fingerprint,
        "source_sheet": full_bundle["source_sheet"],
        "status": "invalid" if errors else "valid_with_warnings" if warnings else "valid",
        "record_count": len(selected),
        "raw_record_count": sum(row["workbook_occurrence_count"] for row in occurrence_validation),
        "numeric_item_code_rows": sum(
            row["workbook_reference"].get("item_code_cell_type") == "n" for row in selected
        ),
        "records": selected,
        "issues": issues,
        "summary": {"errors": errors, "warnings": warnings},
        "scope": {
            "scope_type": "explicit_item_code_allowlist",
            "scope_name": scope_name,
            "selected_item_codes": requested,
            "selected_item_count": len(selected),
            "occurrence_validation": occurrence_validation,
        },
        "source_full_catalog": {
            "catalog_version": full_bundle["catalog_version"],
            "status": full_bundle["status"],
            "summary": full_bundle["summary"],
            "record_count": full_bundle["record_count"],
            "preserved_unchanged": True,
        },
        "known_out_of_scope_issues": known_out_of_scope,
    }


def write_catalog(bundle: dict[str, Any], output_dir: str | Path) -> None:
    output = Path(output_dir)
    if output.exists():
        raise FileExistsError(f"Catalog output already exists; choose a new version directory: {output}")
    output.mkdir(parents=True)
    (output / "product_catalog.json").write_text(
        json.dumps({key: value for key, value in bundle.items() if key != "issues"}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    validation = {
            "schema_version": bundle["schema_version"],
            "catalog_version": bundle["catalog_version"],
            "source_workbook": bundle["source_workbook"],
            "source_workbook_fingerprint": bundle.get("source_workbook_fingerprint"),
            "status": bundle["status"],
            "summary": bundle["summary"],
            "numeric_item_code_rows": bundle["numeric_item_code_rows"],
            "issues": bundle["issues"],
        }
    for key in ("scope", "source_full_catalog", "known_out_of_scope_issues"):
        if key in bundle:
            validation[key] = bundle[key]
    (output / "catalog_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8",
    )
    fields = (
        "item_code", "item_name", "brand", "category", "class_name", "filesystem_name",
        "packaging", "legacy_class_name", "workbook_reference",
    )
    with (output / "product_catalog.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in bundle["records"]:
            row = {name: record.get(name) for name in fields}
            row["workbook_reference"] = json.dumps(row["workbook_reference"], ensure_ascii=False)
            writer.writerow(row)


def load_catalog(path: str | Path) -> dict[str, Any]:
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if value.get("schema_version") != 1 or not isinstance(value.get("records"), list):
        raise ValueError(f"Unsupported catalog schema: {path}")
    codes = [row["item_code"] for row in value["records"]]
    if len(codes) != len(set(codes)):
        raise ValueError("Catalog JSON has duplicate item_code entries")
    return value


def require_usable_catalog(value: dict[str, Any], *, require_scope: bool = False) -> None:
    """Reject globally invalid or unscoped catalogs at demo/runtime boundaries."""

    if value.get("status") not in {"valid", "valid_with_warnings"}:
        raise ValueError(f"Catalog status is not usable: {value.get('status')!r}")
    if require_scope and value.get("scope", {}).get("scope_type") != "explicit_item_code_allowlist":
        raise ValueError("AXE demo requires an explicitly scoped item-code catalog")
    selected = value.get("scope", {}).get("selected_item_codes")
    if require_scope and selected != [row["item_code"] for row in value["records"]]:
        raise ValueError("Scoped catalog record order/content differs from its explicit allowlist")
